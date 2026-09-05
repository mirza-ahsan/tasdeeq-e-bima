"""Build the pitch deck as both .pptx and .pdf from one content source.

    uv run python docs/build_deck.py

The slides are defined once in SLIDES below. The PPTX is written with python-pptx; the
PDF is produced by rendering an HTML version through headless Chromium, so the two never
drift apart in content.
"""

from __future__ import annotations

import html
import shutil
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

DOCS = Path(__file__).resolve().parent

INK = RGBColor(0x17, 0x15, 0x0F)
INK_SOFT = RGBColor(0x4A, 0x46, 0x3B)
MUTED = RGBColor(0x7D, 0x77, 0x67)
ACCENT = RGBColor(0x1F, 0x4B, 0x43)
PAPER = RGBColor(0xFA, 0xF9, 0xF6)
RULE = RGBColor(0xE2, 0xDE, 0xD2)

# (eyebrow, headline, [body lines], [stat (value, label)], note)
SLIDES: list[dict] = [
    dict(eyebrow="BANO QABIL AI HACKATHON", headline="Claim Check",
         body=["A spell-checker for insurance claims.",
               "Checks a clinic's claim before it is submitted: how likely it is to be",
               "rejected, which field is the problem, and what to fix."],
         note="Mirza Ahsan Baig and Safee · LightGBM + SHAP + Qwen on DashScope"),

    dict(eyebrow="THE PROBLEM", headline="Rejections are paperwork, not medicine",
         body=["A clinic treats an insured patient, then bills the insurer or TPA.",
               "A large share of those claims come back rejected — for missing",
               "pre-authorisation, a filing deadline passed, documents left out.",
               "",
               "The clinic finds out weeks later. Small clinics have no billing",
               "specialist to catch it. Cash flow becomes unpredictable, and clinics",
               "eventually stop accepting insurance patients at all."],
         note="A financial inclusion problem as much as a healthcare one"),

    dict(eyebrow="WHAT IT DOES", headline="Three parts, none of them mocked",
         body=["1.  Risk model — LightGBM gives a calibrated probability; SHAP names",
               "     the specific field driving it; a second model predicts the real",
               "     X12 CARC reason code.",
               "",
               "2.  Adaptive questioning — after every answer it computes the mutual",
               "     information between the outcome and each unanswered field, then",
               "     asks whichever would reduce uncertainty most.",
               "",
               "3.  Plain language — Qwen restates the computed result in two",
               "     sentences. It cannot change the number, the code, or the fields."],
         note=""),

    dict(eyebrow="THE DIFFERENTIATOR", headline="It asks, it doesn't just collect",
         body=["Most tools in this space are a form with a model bolted on.",
               "Insurance chatbots that do adaptive intake run a decision tree.",
               "",
               "Ours is information-theoretic: the next question is the one with the",
               "highest mutual information with the outcome, given what is known.",
               "",
               "Across 200 held-out claims we see 72 distinct question orders."],
         stats=[("72", "distinct question orders"), ("3.7", "questions, clean claim"),
                ("4.1", "questions, problem claim")],
         note=""),

    dict(eyebrow="RESULTS", headline="Deliberately not perfect",
         body=["Only 26.5% of claims are rejected, so blindly approving everything",
               "already scores 73.5%. Accuracy is the wrong headline — AUC and recall",
               "are the meaningful numbers.",
               "",
               "Per-code accuracy is highest where it matters: 92% on CARC 197,",
               "pre-authorisation absent."],
         stats=[("0.788", "ROC AUC"), ("67.7%", "recall"), ("68 / 80%", "CARC top-1 / top-3")],
         note="Held at 77.9% accuracy on purpose — see next slide"),

    dict(eyebrow="THE DATA", headline="No, the data is not real. It cannot be.",
         body=["No public dataset exists anywhere of real clinic claims paired with",
               "their outcomes. Insurers treat that as private business information.",
               "Every commercial product here trains on its own private history.",
               "",
               "So we built ours from two real pieces:",
               "     Synthea — open clinical simulator, real SNOMED CT codes",
               "     X12 CARC — the real industry rejection-reason standard",
               "",
               "Only the labelling logic is ours, and we documented exactly where",
               "that line falls. A near-perfect score on self-generated data would",
               "only prove the model memorised our own rules."],
         note="docs/data-provenance.md · docs/model-card.md"),

    dict(eyebrow="THE LOOP", headline="Honest about self-improvement",
         body=["When staff disagree, one click logs the correction as a labelled",
               "example. That is all it does, and we say so.",
               "",
               "Real products close this loop by parsing the 835 remittance files",
               "insurers return, retraining as outcomes accumulate over weeks.",
               "Nobody in this industry has live self-improvement.",
               "",
               "Our architecture supports that path. The demo does not pretend to",
               "have walked it."],
         note=""),

    dict(eyebrow="WHAT'S NEXT", headline="What production would need",
         body=["1.  Real historical claims with real adjudication outcomes, from a",
               "     partner clinic or TPA",
               "2.  Parsed 835 remittance advice, for the codes insurers actually sent",
               "3.  Retraining as outcomes accumulate, with human corrections as labels",
               "4.  Per-insurer models — filing rules and tariffs differ by payer",
               "",
               "The system supports all four. It cannot access the data they need."],
         note="Python · LightGBM · SHAP · FastAPI · Next.js · Qwen on DashScope · Alibaba Cloud ECS"),
]


def _txt(frame, text, size, color, bold=False, font="IBM Plex Sans", space_after=0):
    frame.text = text
    p = frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.space_after = Pt(space_after)
    r = p.runs[0]
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.name = font
    return p


def build_pptx(path: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    for i, s in enumerate(SLIDES):
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = PAPER

        _txt(slide.shapes.add_textbox(Inches(0.9), Inches(0.65), Inches(11), Inches(0.3))
             .text_frame, s["eyebrow"], 11, MUTED, font="IBM Plex Mono")

        _txt(slide.shapes.add_textbox(Inches(0.85), Inches(1.05), Inches(11.6), Inches(1.0))
             .text_frame, s["headline"], 38 if i == 0 else 30, INK, font="Newsreader")

        line = slide.shapes.add_connector(1, Inches(0.9), Inches(2.15), Inches(12.4), Inches(2.15))
        line.line.color.rgb = RULE
        line.line.width = Emu(9525)

        body = s.get("body", [])
        if body:
            tb = slide.shapes.add_textbox(Inches(0.9), Inches(2.5), Inches(7.6), Inches(3.6))
            tf = tb.text_frame
            tf.word_wrap = True
            for j, ln in enumerate(body):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.text = ln
                p.space_after = Pt(3)
                if p.runs:
                    r = p.runs[0]
                    r.font.size = Pt(15)
                    r.font.color.rgb = INK_SOFT
                    r.font.name = "IBM Plex Sans"

        for k, (value, label) in enumerate(s.get("stats", [])):
            top = Inches(2.6 + k * 1.35)
            _txt(slide.shapes.add_textbox(Inches(9.0), top, Inches(3.4), Inches(0.7)).text_frame,
                 value, 34, ACCENT, font="Newsreader")
            _txt(slide.shapes.add_textbox(Inches(9.05), top + Inches(0.62), Inches(3.4), Inches(0.4))
                 .text_frame, label, 11, MUTED, font="IBM Plex Mono")

        if s.get("note"):
            _txt(slide.shapes.add_textbox(Inches(0.9), Inches(6.6), Inches(11.5), Inches(0.4))
                 .text_frame, s["note"], 11, MUTED, font="IBM Plex Mono")

        _txt(slide.shapes.add_textbox(Inches(12.3), Inches(6.6), Inches(0.6), Inches(0.4))
             .text_frame, str(i + 1), 11, MUTED, font="IBM Plex Mono")

    prs.save(path)


def build_html(path: Path) -> None:
    """Same content, laid out for print. Chromium turns this into the PDF."""
    slides = []
    for i, s in enumerate(SLIDES):
        stats = "".join(
            f'<div class="stat"><div class="sv">{html.escape(v)}</div>'
            f'<div class="sl">{html.escape(l)}</div></div>'
            for v, l in s.get("stats", []))
        body = "".join(
            f'<p class="{"gap" if not ln else ""}">{html.escape(ln) or "&nbsp;"}</p>'
            for ln in s.get("body", []))
        slides.append(f"""
<section class="slide">
  <div class="eyebrow">{html.escape(s['eyebrow'])}</div>
  <h1 class="{'lead' if i == 0 else ''}">{html.escape(s['headline'])}</h1>
  <hr/>
  <div class="cols"><div class="body">{body}</div><div class="stats">{stats}</div></div>
  <div class="foot"><span>{html.escape(s.get('note', ''))}</span><span>{i + 1}</span></div>
</section>""")

    path.write_text(f"""<!doctype html><meta charset="utf-8">
<title>Claim Check — deck</title>
<style>
  @page {{ size: 13.333in 7.5in; margin: 0; }}
  * {{ box-sizing: border-box; }}
  body {{ margin: 0; font-family: 'IBM Plex Sans', Helvetica, Arial, sans-serif;
          color: #17150f; background: #faf9f6; }}
  .slide {{ width: 13.333in; height: 7.5in; padding: 0.65in 0.9in; position: relative;
            page-break-after: always; display: flex; flex-direction: column; background: #faf9f6; }}
  .eyebrow {{ font-family: 'IBM Plex Mono', monospace; font-size: 11pt; letter-spacing: .14em;
              color: #7d7767; }}
  h1 {{ font-family: Newsreader, Georgia, serif; font-weight: 400; font-size: 30pt;
        margin: .18in 0 0; letter-spacing: -.01em; }}
  h1.lead {{ font-size: 40pt; }}
  hr {{ border: 0; border-top: 1px solid #e2ded2; margin: .28in 0 .3in; }}
  .cols {{ display: flex; gap: .7in; flex: 1; }}
  .body {{ flex: 1; }}
  .body p {{ font-size: 15pt; line-height: 1.5; color: #4a463b; margin: 0 0 3pt; }}
  .body p.gap {{ height: 9pt; margin: 0; }}
  .stats {{ width: 3.2in; }}
  .stat {{ margin-bottom: .42in; }}
  .sv {{ font-family: Newsreader, Georgia, serif; font-size: 34pt; color: #1f4b43; line-height: 1; }}
  .sl {{ font-family: 'IBM Plex Mono', monospace; font-size: 10.5pt; color: #7d7767; margin-top: 5pt; }}
  .foot {{ display: flex; justify-content: space-between; font-family: 'IBM Plex Mono', monospace;
           font-size: 10.5pt; color: #7d7767; }}
</style>
{''.join(slides)}""", encoding="utf-8")


def main() -> int:
    pptx_path, html_path, pdf_path = DOCS / "deck.pptx", DOCS / "deck.html", DOCS / "deck.pdf"
    build_pptx(pptx_path)
    build_html(html_path)
    print(f"  {pptx_path.name}: {pptx_path.stat().st_size // 1024} KB, {len(SLIDES)} slides")

    chrome = next((c for c in ("chromium", "chromium-browser", "google-chrome")
                   if shutil.which(c)), None)
    if not chrome:
        print("  chromium not found — deck.pdf not built; open deck.html and print to PDF")
        return 0

    subprocess.run([chrome, "--headless", "--disable-gpu", "--no-sandbox",
                    "--no-pdf-header-footer", f"--print-to-pdf={pdf_path}",
                    html_path.as_uri()], check=True, capture_output=True, timeout=120)
    print(f"  {pdf_path.name}:  {pdf_path.stat().st_size // 1024} KB")
    return 0


if __name__ == "__main__":
    sys.exit(main())
